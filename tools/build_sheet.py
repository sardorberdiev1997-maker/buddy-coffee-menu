"""Собирает menu.xlsx из data/menu.json — для первичной загрузки в Google Таблицы.

Запуск: python3 tools/build_sheet.py  →  tools/menu.xlsx
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = json.load(open(os.path.join(ROOT, 'data', 'menu.json'), encoding='utf-8'))

RED = PatternFill('solid', fgColor='E31B23')
HEAD_FONT = Font(bold=True, color='FFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')

WIDTHS = {
    'Меню': [12, 18, 18, 42, 42, 22, 22, 30, 20, 12, 11],
    'Категории': [12, 18, 18, 9],
    'Настройки': [22, 48, 48],
}

# Колонки, которые должны остаться текстом (иначе Google превратит «15000 / 20000» и «15000» в разные типы)
TEXT_COLS = {'Меню': {'цена', 'размеры_uz', 'размеры_ru', 'фото', 'метка', 'в_наличии'}}

INSTRUCTION = [
    ['Как редактировать меню'],
    [''],
    ['Сайт сам подтягивает изменения из этой таблицы — обычно в течение минуты. Достаточно обновить страницу меню.'],
    ['Не переименовывайте вкладки «Меню», «Категории», «Настройки» и не меняйте первую строку с названиями колонок.'],
    [''],
    ['Вкладка «Меню» — одна строка = одна позиция. Порядок строк = порядок на сайте.'],
    ['категория — id категории из вкладки «Категории» (hotdog, lavash, shaurma, kfs, burger, ichimlik).'],
    ['название_uz / название_ru — название на узбекском и русском.'],
    ['описание_uz / описание_ru — короткое описание (состав, вес, объём).'],
    ['цена — число без пробелов: 35000. Если размеров несколько — через косую черту: 15000 / 20000 / 25000.'],
    ['размеры_uz / размеры_ru — названия размеров в том же порядке, что и цены: Kichik / O‘rta / Katta. Для одной цены оставьте пустым.'],
    ['фото — имя файла из папки img на сайте (например lavash) ИЛИ прямая ссылка на картинку (https://...). Ссылка на файл Google Диска тоже подойдёт, если у файла включён «Доступ по ссылке».'],
    ['метка — хит, новинка, острое или советуем. Можно оставить пустым.'],
    ['в_наличии — «да» или «нет». «нет» — позиция станет серой с пометкой «Нет в наличии» (удалять строку не нужно).'],
    [''],
    ['Вкладка «Категории» — порядок строк = порядок разделов на сайте. Иконка — любой эмодзи.'],
    ['Чтобы добавить категорию: новая строка с латинским id без пробелов, затем используйте этот id в колонке «категория» на вкладке «Меню».'],
    [''],
    ['Вкладка «Настройки» — тексты шапки и блока контактов. Пустое значение = строка не показывается на сайте.'],
    ['ссылка_на_карту — ссылка «Поделиться» из Google Maps или Яндекс Карт.'],
    ['wifi_сеть / wifi_пароль — если заполнить, гости увидят пароль и кнопку «Копировать».'],
    ['объявление — жёлтая плашка под шапкой (акция, «сегодня до 20:00» и т.п.).'],
    [''],
    ['Если что-то сломалось — сайт покажет последнюю удачно загруженную версию меню.'],
]


def fill_sheet(ws, rows, name):
    text_cols = TEXT_COLS.get(name, set())
    header = rows[0]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = WRAP
            if header[c - 1] in text_cols:
                cell.number_format = '@'
            if r == 1:
                cell.fill = RED
                cell.font = HEAD_FONT
    # Формат «текст» и для пустых строк ниже — чтобы новые цены тоже были текстом
    for c, h in enumerate(header, start=1):
        if h in text_cols:
            for r in range(len(rows) + 1, 101):
                ws.cell(row=r, column=c).number_format = '@'
    for i, w in enumerate(WIDTHS.get(name, []), start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = 'A2'


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for name in ('Меню', 'Категории', 'Настройки'):
        ws = wb.create_sheet(name)
        fill_sheet(ws, DATA[name], name)

    menu = wb['Меню']
    head = DATA['Меню'][0]
    col = lambda key: menu.cell(row=1, column=head.index(key) + 1).column_letter

    dv_badge = DataValidation(type='list', formula1='"хит,новинка,острое,советуем"', allow_blank=True)
    dv_avail = DataValidation(type='list', formula1='"да,нет"', allow_blank=True)
    dv_cat = DataValidation(type='list', formula1="='Категории'!$A$2:$A$30", allow_blank=False)
    for dv, key in ((dv_badge, 'метка'), (dv_avail, 'в_наличии'), (dv_cat, 'категория')):
        menu.add_data_validation(dv)
        dv.add(f'{col(key)}2:{col(key)}100')

    ws = wb.create_sheet('Инструкция')
    for r, (line,) in enumerate(INSTRUCTION, start=1):
        cell = ws.cell(row=r, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 120

    out = os.path.join(ROOT, 'tools', 'menu.xlsx')
    wb.save(out)
    print(out)


if __name__ == '__main__':
    main()
